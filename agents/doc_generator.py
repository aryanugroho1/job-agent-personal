import os
import re
import subprocess
from pathlib import Path
from docxtpl import DocxTemplate
from playwright.sync_api import sync_playwright

TEMPLATES_DIR = Path("templates")
GENERATED_DIR = Path("generated_cvs")

def ensure_directories():
    TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)
    GENERATED_DIR.mkdir(parents=True, exist_ok=True)

def clean_md(text: str) -> str:
    """Replaces markdown bold **text** with HTML <strong>text</strong>."""
    return re.sub(r"\*\*(.*?)\*\*", r"<strong>\1</strong>", text)

def render_pdf_playwright(master_cv: dict, tailored_title: str, tailored_summary: str, out_pdf_path: Path):
    """
    Renders an executive, ATS-friendly single/two-page resume directly to vector PDF using Playwright Chromium.
    """
    candidate = master_cv["candidate"]
    competencies = master_cv.get("competencies", {})
    experience = master_cv.get("experience", [])
    education = master_cv.get("education", [])
    certifications = master_cv.get("certifications_awards", [])

    html_content = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<style>
  @page {{
    size: A4;
    margin: 14mm 14mm 14mm 14mm;
  }}
  body {{
    font-family: 'Segoe UI', -apple-system, BlinkMacSystemFont, Roboto, Helvetica, Arial, sans-serif;
    color: #1e293b;
    line-height: 1.45;
    font-size: 10pt;
    margin: 0;
    padding: 15px 25px;
    background: #ffffff;
  }}
  .header {{
    border-bottom: 2px solid #0284c7;
    padding-bottom: 8px;
    margin-bottom: 12px;
  }}
  .name {{
    font-size: 19pt;
    font-weight: 700;
    color: #0f172a;
    letter-spacing: -0.5px;
    margin: 0;
  }}
  .target-title {{
    font-size: 12.5pt;
    font-weight: 600;
    color: #0284c7;
    margin: 3px 0 5px 0;
  }}
  .contacts {{
    font-size: 9pt;
    color: #475569;
  }}
  .contacts a {{
    color: #0284c7;
    text-decoration: none;
  }}
  h2 {{
    font-size: 10.5pt;
    text-transform: uppercase;
    letter-spacing: 0.8px;
    color: #0f172a;
    border-bottom: 1px solid #cbd5e1;
    padding-bottom: 2px;
    margin-top: 12px;
    margin-bottom: 6px;
  }}
  p.summary {{
    margin: 4px 0 8px 0;
    text-align: justify;
    color: #334155;
    font-size: 9.5pt;
  }}
  .competency-grid {{
    display: grid;
    grid-template-columns: 1fr;
    gap: 3px;
    margin-bottom: 8px;
  }}
  .comp-item {{
    font-size: 9pt;
    color: #334155;
  }}
  .comp-category {{
    font-weight: 600;
    color: #0f172a;
  }}
  .exp-item {{
    margin-bottom: 10px;
  }}
  .exp-header {{
    display: flex;
    justify-content: space-between;
    font-weight: 700;
    color: #0f172a;
    font-size: 10pt;
  }}
  .exp-sub {{
    display: flex;
    justify-content: space-between;
    font-style: italic;
    color: #0369a1;
    font-size: 9pt;
    margin-bottom: 3px;
  }}
  ul {{
    margin: 2px 0 4px 16px;
    padding: 0;
  }}
  li {{
    margin-bottom: 2px;
    color: #334155;
    font-size: 9pt;
  }}
  .edu-item, .cert-item {{
    font-size: 9pt;
    margin-bottom: 3px;
    color: #334155;
  }}
</style>
</head>
<body>

<div class="header">
  <div class="name">{candidate['name']}</div>
  <div class="target-title">{tailored_title}</div>
  <div class="contacts">
    {candidate['location']} &nbsp;|&nbsp; 
    {candidate['phone']} &nbsp;|&nbsp; 
    <a href="mailto:{candidate['email']}">{candidate['email']}</a> &nbsp;|&nbsp; 
    <a href="{candidate['linkedin']}">{candidate['linkedin']}</a>
  </div>
</div>

<h2>Professional Summary</h2>
<p class="summary">{tailored_summary}</p>

<h2>Core Competencies & Technical Skills</h2>
<div class="competency-grid">
"""
    for cat, skills in competencies.items():
        skills_str = ", ".join(skills).rstrip(".")
        html_content += f"""  <div class="comp-item"><span class="comp-category">• {cat}:</span> {skills_str}</div>\n"""

    html_content += """</div>

<h2>Professional Experience</h2>
"""
    for exp in experience[:6]:
        html_content += f"""
<div class="exp-item">
  <div class="exp-header">
    <span>{exp['company']}</span>
    <span>{exp.get('location', '')}</span>
  </div>
  <div class="exp-sub">
    <span>{exp['role']}</span>
    <span>{exp['period']}</span>
  </div>
  <ul>
"""
        for h in exp["highlights"]:
            html_content += f"    <li>{clean_md(h)}</li>\n"
        html_content += """  </ul>
</div>
"""

    html_content += """
<h2>Education</h2>
"""
    for edu in education:
        html_content += f"""<div class="edu-item"><strong>{edu['degree']}</strong> — {edu['institution']}</div>\n"""

    html_content += """
<h2>Certifications & Honors</h2>
<ul>
"""
    for cert in certifications[:6]:
        html_content += f"""  <li class="cert-item">{clean_md(cert)}</li>\n"""

    html_content += """</ul>

</body>
</html>
"""

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        page = browser.new_page()
        page.set_content(html_content, wait_until="load")
        page.pdf(
            path=str(out_pdf_path),
            format="A4",
            print_background=True,
            margin={"top": "10mm", "bottom": "10mm", "left": "10mm", "right": "10mm"}
        )
        browser.close()


def generate_custom_cv(
    master_cv: dict,
    tailored_title: str,
    tailored_summary: str,
    company_name: str,
    job_id: str,
    template_name: str = "cv_template.docx"
) -> Path:
    """
    Renders tailored resume into an executive PDF using Playwright Chromium.
    Also produces DOCX copy as fallback.
    """
    ensure_directories()
    safe_company = "".join(c for c in company_name if c.isalnum() or c in (" ", "_", "-")).strip().replace(" ", "_")
    out_pdf_path = GENERATED_DIR / f"{safe_company}_{job_id}.pdf"
    out_docx_path = GENERATED_DIR / f"{safe_company}_{job_id}.docx"

    # 1. Primary: Generate clean vector PDF via Playwright
    try:
        render_pdf_playwright(master_cv, tailored_title, tailored_summary, out_pdf_path)
        print(f"[Doc Generator] Generated tailored PDF: {out_pdf_path}")
        return out_pdf_path
    except Exception as e:
        print(f"[Doc Generator Error] Playwright PDF generation failed: {e}")

    # 2. Fallback: DOCX + LibreOffice
    template_path = TEMPLATES_DIR / template_name
    if not template_path.exists():
        from docx import Document
        doc = Document()
        doc.add_heading("ARYA MAULANA NUGROHO, ITIL®", level=1)
        doc.add_paragraph("{{ target_role_title }}")
        doc.add_paragraph("Tokyo, Japan | +817090934764 | arya.nugroho.works@gmail.com | linkedin.com/in/aryanugroho1/")
        doc.add_heading("Professional Summary", level=2)
        doc.add_paragraph("{{ tailored_summary }}")
        doc.add_heading("Experience", level=2)
        doc.add_paragraph("{% for exp in experience %}")
        doc.add_paragraph("{{ exp.company }} - {{ exp.role }} ({{ exp.period }})")
        doc.add_paragraph("{% for h in exp.highlights %}")
        doc.add_paragraph("• {{ h }}")
        doc.add_paragraph("{% endfor %}")
        doc.add_paragraph("{% endfor %}")
        doc.save(template_path)

    doc = DocxTemplate(template_path)
    context = {
        "candidate_name": master_cv["candidate"]["name"],
        "target_role_title": tailored_title,
        "tailored_summary": tailored_summary,
        "candidate": master_cv["candidate"],
        "competencies": master_cv.get("competencies", {}),
        "experience": master_cv.get("experience", []),
        "education": master_cv.get("education", []),
        "certifications_awards": master_cv.get("certifications_awards", [])
    }
    doc.render(context)
    doc.save(out_docx_path)

    convert_docx_to_pdf(out_docx_path, GENERATED_DIR)
    return out_pdf_path if out_pdf_path.exists() else out_docx_path


def convert_docx_to_pdf(docx_path: Path, output_dir: Path):
    """Converts docx to PDF via libreoffice headless or soffice."""
    for cmd_name in ["libreoffice", "soffice"]:
        try:
            cmd = [cmd_name, "--headless", "--convert-to", "pdf", "--outdir", str(output_dir), str(docx_path)]
            res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=60)
            if res.returncode == 0:
                return
        except Exception:
            pass
    print(f"[Doc Generator Warning] LibreOffice CLI not available for {docx_path}.")
